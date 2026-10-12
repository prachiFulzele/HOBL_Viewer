from __future__ import annotations

import csv
import re
from collections import defaultdict
from datetime import date
from pathlib import Path
from statistics import median

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(r"C:\Users\pfulzele\Downloads\AUTOBUGGING")
REPORT_DATE = date.today().isoformat()

SOURCES = {
    "Cadmus": ROOT / "Cadmus" / "Cadmus_latest.xlsx",
    "LNL": ROOT / "LNL" / "LNL.xlsx",
    "PTL PUE": ROOT / "PTL PUE" / "HOBL_specified_2026-08-28 (1).xlsx",
    "PTL32 GB": ROOT / "PTL32 GB" / "Pantherlake32GB.xlsx",
    "Purwa DC": ROOT / "Purwa DC" / "HOBL_specified_2026-08-24.xlsx",
    "Purwa SC": ROOT / "Purwa SC-mactching system" / "HOBL_specified_2026-08-24.xlsx",
    "ROT": ROOT / "ROT" / "ROT_result.xlsx",
    "WCL": ROOT / "WCL" / "WCL.xlsx",
}

GOAL_WORKBOOKS = {
    "Cadmus": ROOT / "Cadmus" / "Cadmus_Goal_vs_P50_All_Scenarios_2026-08-27.xlsx",
    "PTL PUE": ROOT / "PTL PUE" / "PTL PUE_Goal_vs_P50_All_Scenarios_2026-08-27.xlsx",
    "PTL32 GB": ROOT / "PTL32 GB" / "PTL32 GB_Goal_vs_P50_All_Scenarios_2026-08-27.xlsx",
    "Purwa DC": ROOT / "Purwa DC" / "Purwa DC_Goal_vs_P50_All_Scenarios_2026-08-27.xlsx",
    "Purwa SC": ROOT / "Purwa SC-mactching system" / "Purwa SC_Goal_vs_P50_All_Scenarios_2026-08-27.xlsx",
}
ROT_GOALS = ROOT / "ROT" / "ROT_15_EV_goals.csv"
PERFORMANCE_REFERENCE = ROOT / "All_5_Devices_Performance_Goals_P50_Plus10_Recalculated_2026-08-28.xlsx"

HOBL_OUTPUT = ROOT / f"All_Platforms_HOBL_Goal_vs_P50_{REPORT_DATE}.xlsx"
PERFORMANCE_OUTPUT = ROOT / f"All_8_Devices_Performance_P50_Plus10_{REPORT_DATE}.xlsx"

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF2F8"
WHITE = "FFFFFF"
TEXT = "1F2937"
GREEN = "C6EFCE"
AMBER = "FFEB9C"
RED = "FFC7CE"
GRAY = "E7E6E6"
THIN_GRAY = Side(style="thin", color="D9E1F2")
TABLE_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)

RAW_POWER_RAILS = {
    "pmemisoc": "soc",
    "pmeminpu": "npu",
    "pmemimemory": "memory",
    "pmemi3p3vssd": "storage",
    "pmemi3p86vwifi": "wifi",
    "pmemi3p3vpanelin": "display_logic",
    "pmemibkltin": "display_light",
    "pmemivsysrop": "rop",
    "pmemisystem": "total",
}


def normalized(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def scenario_key(value: object) -> str:
    key = normalized(value)
    aliases = {
        "modernstandbyfloor": "cs_floor",
        "csfloor": "cs_floor",
        "localvideoplayback": "lvp",
        "ablactiveidle": "idle_apps",
        "idleapps": "idle_apps",
        "ablactiveproductivity": "productivity",
        "ablactiveprod": "productivity",
        "ablactiveweb": "web",
        "teams3x3video": "teams2_3x3_video",
        "teams2_3x3video": "teams2_3x3_video",
        "teams11audio": "teams2_1on1_audio",
        "teams2_1on1audio": "teams2_1on1_audio",
        "teams11video": "teams2_1on1_video",
        "teams2_1on1video": "teams2_1on1_video",
        "teams2idle": "teams2_idle",
    }
    return aliases.get(key, re.sub(r"[^a-z0-9]+", "_", str(value or "").lower()).strip("_"))


def power_metric_key(value: object) -> str:
    key = normalized(value)
    aliases = {
        "socpowerw": "soc",
        "memorypowerw": "memory",
        "drampowerw": "memory",
        "ainpupowerw": "npu",
        "npupowerw": "npu",
        "storagepowerw": "storage",
        "ssdpowerw": "storage",
        "wifibtpowerw": "wifi",
        "wifipowerw": "wifi",
        "displaypowerw": "display",
        "lcddisplaypowerw": "display",
        "displaylightpowerw": "display_light",
        "lcdbacklightpowerw": "display_light",
        "displaylogicpowerw": "display_logic",
        "lcdpanelpowerw": "display_logic",
        "touchpowerw": "touch",
        "audiopowerw": "audio",
        "camerapowerw": "camera",
        "sampowerw": "sam",
        "kbtppowerw": "keyboard_trackpad",
        "keyboardtouchpadw": "keyboard_trackpad",
        "keyboardtrackpadpowerw": "keyboard_trackpad",
        "trackpadhapticpowerw": "keyboard_trackpad",
        "roppowerw": "rop",
        "systempowerw": "total",
        "totalpowerw": "total",
    }
    return aliases.get(key, key)


def numeric_values(values, include_zero: bool) -> list[float]:
    result = []
    for value in values:
        if isinstance(value, (int, float)) and (include_zero or value != 0):
            result.append(float(value))
    return result


def performance_cell_values(value) -> list[float]:
    if isinstance(value, (int, float)):
        return [float(value)] if value != 0 else []
    if not isinstance(value, str) or value.strip().upper() == "N/A":
        return []
    result = []
    for part in value.split(","):
        try:
            number = float(part.strip())
        except ValueError:
            continue
        if number != 0:
            result.append(number)
    return result


def percentile(values: list[float], percent: int) -> float:
    ordered = sorted(values)
    position = (len(ordered) - 1) * percent / 100
    lower = int(position)
    upper = min(lower + 1, len(ordered) - 1)
    fraction = position - lower
    return ordered[lower] + (ordered[upper] - ordered[lower]) * fraction


def raw_performance_metric(label: object):
    text = str(label or "")
    match = re.search(r"\(pt\s+(\d+)\)", text, re.IGNORECASE)
    perftrack_id = int(match.group(1)) if match else None
    simple = {
        8998: "edgelaunchfirstlaunch",
        9412: "edgepageload",
        8804: "officeoutlookbootv2",
        8805: "officewordbootv2",
        8806: "officexlbootv2",
        8807: "officepptbootv2",
        9475: "fileexplorerlaunch",
        9971: "snippingtooloverlaylaunchperformance",
        6352: "explorerinputdelay",
    }
    percentile_metrics = {
        1809: "startlaunch",
        9917: "startlaunch",
        10010: "typetosearch",
        2430: "typetosearch",
        9232: "searchfullactivation",
    }
    if perftrack_id in simple:
        return simple[perftrack_id], False
    if perftrack_id in percentile_metrics:
        return percentile_metrics[perftrack_id], True
    return normalized(label), False


def extract_results(path: Path):
    power = defaultdict(lambda: defaultdict(list))
    raw_power = defaultdict(lambda: defaultdict(list))
    performance = defaultdict(lambda: defaultdict(list))
    scenario_labels = {}
    workbook = load_workbook(path, read_only=True, data_only=True)

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 16:
            continue
        scenario_starts = [index for index, value in enumerate(rows[11][1:], 1) if value not in (None, "")]
        for group_index, start in enumerate(scenario_starts):
            raw_scenario = rows[11][start]
            scenario = scenario_key(raw_scenario)
            scenario_labels.setdefault(scenario, str(raw_scenario))
            end = scenario_starts[group_index + 1] if group_index + 1 < len(scenario_starts) else len(rows[14])
            p50_column = next(
                (column for column in range(start, end) if str(rows[14][column]).strip().upper() == "P50"),
                None,
            )
            if p50_column is None:
                continue

            section = None
            raw_rails = {}
            for row in rows[15:]:
                label = row[0] if row else None
                if label == "Power Metrics":
                    section = "power"
                    continue
                if label == "Performance Metrics":
                    section = "performance"
                    continue
                if label == "Power Metrics (raw rails)":
                    section = "raw_power"
                    continue
                if not label:
                    continue
                values = row[start:p50_column]
                if section == "power":
                    power[scenario][power_metric_key(label)].extend(numeric_values(values, include_zero=True))
                elif section == "raw_power":
                    samples = [float(value) if isinstance(value, (int, float)) else None for value in values]
                    raw_rails[normalized(label)] = samples
                    metric = RAW_POWER_RAILS.get(normalized(label))
                    if metric:
                        raw_power[scenario][metric].extend(value for value in samples if value is not None)
                elif section == "performance":
                    metric, add_percentiles = raw_performance_metric(label)
                    cell_samples = [performance_cell_values(value) for value in values]
                    performance[scenario][metric].extend(number for samples in cell_samples for number in samples)
                    if add_percentiles:
                        for percent in (50, 70, 90):
                            key = f"{metric}p{percent}"
                            performance[scenario][key].extend(
                                percentile(samples, percent) for samples in cell_samples if samples
                            )

            panel = raw_rails.get("pmemi3p3vpanelin", [])
            backlight = raw_rails.get("pmemibkltin", [])
            raw_power[scenario]["display"].extend(
                panel_value + backlight_value
                for panel_value, backlight_value in zip(panel, backlight)
                if panel_value is not None and backlight_value is not None
            )

    for scenario, metrics in raw_power.items():
        for metric, samples in metrics.items():
            if not power[scenario].get(metric):
                power[scenario][metric].extend(samples)

    return power, performance, scenario_labels


def load_goal_workbook(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    scenarios = []
    for column in range(2, sheet.max_column + 1, 2):
        label = sheet.cell(3, column).value
        if label:
            scenarios.append((str(label), scenario_key(label), column))

    metrics = []
    targets = defaultdict(dict)
    for row in range(5, sheet.max_row + 1):
        label = sheet.cell(row, 1).value
        if not label or label == "Legend":
            continue
        label = str(label)
        metrics.append((label, power_metric_key(label)))
        for _, scenario, column in scenarios:
            value = sheet.cell(row, column).value
            if isinstance(value, (int, float)):
                targets[scenario][power_metric_key(label)] = float(value)
    return scenarios, metrics, targets


def load_rot_goals(path: Path):
    scenarios = []
    metrics = []
    targets = defaultdict(dict)
    with path.open(newline="", encoding="utf-8-sig") as stream:
        rows = list(csv.reader(stream))
    headers = rows[0]
    display_names = {
        "cs_floor": "Modern Standby Floor",
        "lvp": "Local Video Playback",
        "idle_desktop": "Idle Desktop",
        "abl_active": "ABL Active",
        "abl_standby": "ABL Standby",
        "idle_apps": "ABL Active Idle",
        "productivity": "ABL Active Productivity",
        "web": "Web",
        "teams2_1on1_audio": "Teams 1:1 Audio",
        "teams2_3x3_video": "Teams 3x3 Video",
        "teams2_1on1_video": "Teams 1:1 Video",
        "halo2": "Halo 2",
        "netflix": "Netflix",
        "youtube": "YouTube",
    }
    for column, label in enumerate(headers[1:], 1):
        if label:
            key = scenario_key(label)
            scenarios.append((display_names.get(key, label), key, column))
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        label = row[0]
        metric = power_metric_key(label)
        metrics.append((label, metric))
        for _, scenario, column in scenarios:
            try:
                targets[scenario][metric] = float(row[column])
            except (IndexError, TypeError, ValueError):
                pass
    return scenarios, metrics, targets


def style_title(sheet, title: str, subtitle: str, end_column: int):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    sheet.cell(1, 1, title)
    sheet.cell(2, 1, subtitle)
    sheet.cell(1, 1).font = Font(name="Aptos Display", size=16, bold=True, color=WHITE)
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    sheet.cell(1, 1).alignment = Alignment(vertical="center")
    sheet.cell(2, 1).font = Font(name="Aptos", size=9, color="44546A", italic=True)
    sheet.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 32


def style_table_cell(cell, fill=None, bold=False, color=TEXT, alignment="center"):
    cell.font = Font(name="Aptos", size=10, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill or WHITE)
    cell.border = TABLE_BORDER
    cell.alignment = Alignment(horizontal=alignment, vertical="center", wrap_text=True)


def create_hobl_report(extracted):
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    style_title(
        overview,
        "HOBL Test Plan - Goal vs Latest P50",
        "Power P50 = median of all numeric raw iterations, including zero. Status: Pass <= goal; Watch <= goal +10%; Miss > goal +10%.",
        6,
    )
    headers = ["Platform", "Latest result", "Goal source", "Scenarios", "Goal coverage", "Notes"]
    for column, value in enumerate(headers, 1):
        overview.cell(4, column, value)
        style_table_cell(overview.cell(4, column), BLUE, True, WHITE)

    for row, (platform, source) in enumerate(SOURCES.items(), 5):
        power, _, labels = extracted[platform]
        goal_source = GOAL_WORKBOOKS.get(platform) or (ROT_GOALS if platform == "ROT" else None)
        sample_count = sum(len(samples) for metrics in power.values() for samples in metrics.values())
        if sample_count == 0:
            note = "No numeric power measurements in latest result"
        elif sample_count < 10:
            note = f"Limited power coverage: {sample_count} numeric samples"
        else:
            note = ""
        values = [
            platform,
            source.name,
            goal_source.name if goal_source else "Not available",
            len(labels),
            "Available" if goal_source else "Unavailable",
            note,
        ]
        for column, value in enumerate(values, 1):
            overview.cell(row, column, value)
            style_table_cell(overview.cell(row, column), PALE_BLUE if row % 2 else WHITE, alignment="left" if column in (1, 2, 3, 6) else "center")

    overview.freeze_panes = "A5"
    overview.auto_filter.ref = f"A4:F{overview.max_row}"
    for column, width in enumerate([18, 42, 46, 12, 16, 44], 1):
        overview.column_dimensions[get_column_letter(column)].width = width

    for platform in SOURCES:
        power, _, source_labels = extracted[platform]
        if platform in GOAL_WORKBOOKS:
            scenarios, metrics, targets = load_goal_workbook(GOAL_WORKBOOKS[platform])
            goal_note = GOAL_WORKBOOKS[platform].name
        elif platform == "ROT":
            scenarios, metrics, targets = load_rot_goals(ROT_GOALS)
            goal_note = ROT_GOALS.name
        else:
            scenario_keys = list(source_labels)
            scenarios = [(source_labels[key].replace("_", " ").title(), key, 0) for key in scenario_keys]
            metric_keys = []
            for scenario in scenario_keys:
                for metric in power[scenario]:
                    if metric not in metric_keys:
                        metric_keys.append(metric)
            metrics = [(metric.replace("_", " ").title(), metric) for metric in metric_keys]
            targets = defaultdict(dict)
            goal_note = "No platform goal source found"

        goal_scenarios = {scenario for _, scenario, _ in scenarios}
        scenarios.extend(
            (source_labels[scenario].replace("_", " ").title(), scenario, 0)
            for scenario in source_labels
            if scenario not in goal_scenarios and power[scenario]
        )

        sheet = workbook.create_sheet(platform[:31])
        end_column = 1 + 2 * max(1, len(scenarios))
        style_title(
            sheet,
            f"{platform} Power Goals vs Latest P50",
            f"Result: {SOURCES[platform].name} | Goal source: {goal_note} | N/A indicates missing goal or numeric measurement.",
            end_column,
        )
        sheet.cell(4, 1, "Sub-system")
        style_table_cell(sheet.cell(4, 1), BLUE, True, WHITE, "left")
        sheet.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)
        for index, (label, _, _) in enumerate(scenarios):
            column = 2 + index * 2
            sheet.merge_cells(start_row=4, start_column=column, end_row=4, end_column=column + 1)
            sheet.cell(4, column, label.replace("_", " ").title() if label == label.lower() else label)
            style_table_cell(sheet.cell(4, column), BLUE, True, WHITE)
            sheet.cell(5, column, "Goal")
            sheet.cell(5, column + 1, "P50")
            style_table_cell(sheet.cell(5, column), LIGHT_BLUE, True)
            style_table_cell(sheet.cell(5, column + 1), LIGHT_BLUE, True)

        if not metrics:
            metrics = [("No numeric power metrics in latest result", "__none__")]
        for row_index, (metric_label, metric) in enumerate(metrics, 6):
            sheet.cell(row_index, 1, metric_label)
            style_table_cell(sheet.cell(row_index, 1), PALE_BLUE, True, alignment="left")
            for scenario_index, (_, scenario, _) in enumerate(scenarios):
                goal_column = 2 + scenario_index * 2
                p50_column = goal_column + 1
                goal = targets[scenario].get(metric)
                samples = power[scenario].get(metric, [])
                p50 = median(samples) if samples else None
                sheet.cell(row_index, goal_column, goal if goal is not None else "N/A")
                sheet.cell(row_index, p50_column, p50 if p50 is not None else "N/A")
                style_table_cell(sheet.cell(row_index, goal_column), WHITE)
                style_table_cell(sheet.cell(row_index, p50_column), WHITE)
                if goal is not None and p50 is not None:
                    fill = GREEN if p50 <= goal else (AMBER if p50 <= goal * 1.10 else RED)
                    sheet.cell(row_index, p50_column).fill = PatternFill("solid", fgColor=fill)
                else:
                    sheet.cell(row_index, p50_column).fill = PatternFill("solid", fgColor=GRAY)
                for column in (goal_column, p50_column):
                    if isinstance(sheet.cell(row_index, column).value, (int, float)):
                        sheet.cell(row_index, column).number_format = "0.000"

        sheet.freeze_panes = "B6"
        sheet.sheet_view.showGridLines = False
        sheet.column_dimensions["A"].width = 31
        for column in range(2, end_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 12
        sheet.auto_filter.ref = f"A5:{get_column_letter(end_column)}{sheet.max_row}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    workbook.save(HOBL_OUTPUT)


def performance_metric_order():
    workbook = load_workbook(PERFORMANCE_REFERENCE, read_only=True, data_only=True)
    order = []
    labels = {}
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=5, values_only=True):
            label = row[1] if len(row) > 1 else None
            key = normalized(label)
            if label and key not in labels:
                labels[key] = str(label)
                order.append(key)
    return order, labels


def create_performance_report(extracted):
    metric_order, metric_labels = performance_metric_order()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Performance Thresholds"
    style_title(
        sheet,
        "All Platforms - Performance P50 and +10% Threshold",
        "Consumer combines consumer_multitasker and mincp_base. Enterprise uses enterprise_collab. P50 excludes zero, blank, and N/A; threshold = P50 x 1.10.",
        6,
    )
    headers = ["Test plan", "Platform", "Metric", "P50", "Threshold (+10%)", "Runs"]
    for column, value in enumerate(headers, 1):
        sheet.cell(4, column, value)
        style_table_cell(sheet.cell(4, column), BLUE, True, WHITE)

    workloads = [
        ("Enterprise Collab", {"enterprise_collab"}),
        ("Consumer Multitasker", {"consumer_multitasker", "mincp_base"}),
    ]
    output_row = 5
    for workload_label, scenario_set in workloads:
        for platform in SOURCES:
            _, performance, _ = extracted[platform]
            for metric in metric_order:
                samples = []
                for scenario in scenario_set:
                    samples.extend(performance[scenario].get(metric, []))
                p50 = median(samples) if samples else None
                values = [
                    workload_label,
                    platform,
                    metric_labels[metric],
                    p50 if p50 is not None else "N/A",
                    p50 * 1.10 if p50 is not None else "N/A",
                    len(samples),
                ]
                for column, value in enumerate(values, 1):
                    sheet.cell(output_row, column, value)
                    style_table_cell(
                        sheet.cell(output_row, column),
                        PALE_BLUE if output_row % 2 else WHITE,
                        alignment="left" if column <= 3 else "center",
                    )
                if p50 is not None:
                    sheet.cell(output_row, 4).number_format = "0.0"
                    sheet.cell(output_row, 5).number_format = "0.0"
                output_row += 1

    sheet.freeze_panes = "D5"
    sheet.auto_filter.ref = f"A4:F{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    for column, width in enumerate([23, 17, 43, 14, 20, 10], 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(PERFORMANCE_OUTPUT)


def main():
    missing = [str(path) for path in [*SOURCES.values(), *GOAL_WORKBOOKS.values(), ROT_GOALS, PERFORMANCE_REFERENCE] if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing required inputs:\n" + "\n".join(missing))

    extracted = {platform: extract_results(path) for platform, path in SOURCES.items()}
    create_hobl_report(extracted)
    create_performance_report(extracted)
    print(HOBL_OUTPUT)
    print(PERFORMANCE_OUTPUT)


if __name__ == "__main__":
    main()